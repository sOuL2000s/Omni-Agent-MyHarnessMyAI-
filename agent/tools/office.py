"""Precise Word / Excel / PowerPoint tools. Paragraph-level edits preserve styles."""
from .registry import tool
from ..safety import resolve

S = {"type": "string"}
I = {"type": "integer"}


def _docx():
    try:
        import docx
        return docx
    except ImportError:
        raise ValueError("python-docx missing: pip install python-docx")


@tool("docx_read", "Read a .docx: every paragraph with its index and style, plus tables. Use indexes for edits.",
      {"path": S}, ["path"])
def docx_read(cfg, path):
    d = _docx().Document(str(resolve(cfg, path)))
    out = [f"[{i}] ({p.style.name}) {p.text}" for i, p in enumerate(d.paragraphs)]
    for ti, t in enumerate(d.tables):
        out.append(f"--- table {ti} ({len(t.rows)}x{len(t.columns)}) ---")
        for ri, r in enumerate(t.rows):
            out.append(f"T{ti}R{ri}: " + " | ".join(c.text for c in r.cells))
    return "\n".join(out)


def _set_par_text(p, text):
    """Replace paragraph text keeping the first run's formatting."""
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


@tool("docx_edit_paragraph", "Replace the full text of paragraph `index` (keeps first-run formatting).",
      {"path": S, "index": I, "text": S}, ["path", "index", "text"])
def docx_edit_paragraph(cfg, path, index, text):
    p = resolve(cfg, path, write=True)
    d = _docx().Document(str(p))
    n = len(d.paragraphs)
    if not (0 <= index < n):
        raise ValueError(f"paragraph index {index} out of range (document has {n} paragraphs)")
    _set_par_text(d.paragraphs[index], text)
    d.save(str(p))
    return f"Paragraph {index} updated"


@tool("docx_replace_text", "Find/replace across the document (body, tables). Keeps run formatting where possible.",
      {"path": S, "find": S, "replace": S}, ["path", "find", "replace"])
def docx_replace_text(cfg, path, find, replace):
    p = resolve(cfg, path, write=True)
    d = _docx().Document(str(p))
    pars = list(d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                pars.extend(c.paragraphs)
    n = 0
    for par in pars:
        if find not in par.text:
            continue
        done = False
        for r in par.runs:
            if find in r.text:
                r.text = r.text.replace(find, replace)
                n += 1
                done = True
        if not done:  # text spans several runs
            _set_par_text(par, par.text.replace(find, replace))
            n += 1
    d.save(str(p))
    return f"{n} paragraph(s)/run(s) changed"


@tool("docx_insert_paragraph",
      "Insert a new paragraph after paragraph `after_index` (-1 = at end) with optional style.",
      {"path": S, "after_index": I, "text": S, "style": S}, ["path", "text"])
def docx_insert_paragraph(cfg, path, text, after_index=-1, style=None):
    p = resolve(cfg, path, write=True)
    d = _docx().Document(str(p))
    n = len(d.paragraphs)
    # `after_index` is the index to insert *after*. Appending at the end
    # means "insert after the last paragraph", which is `after_index == n-1`.
    # Any index >= n is out of range and is treated as "at end" (forgiving,
    # since a model can easily overshoot).
    if after_index < 0 or after_index >= n - 1:
        d.add_paragraph(text, style=style)
    else:
        d.paragraphs[after_index + 1].insert_paragraph_before(text, style=style)
    d.save(str(p))
    return f"inserted after paragraph {after_index if after_index >= 0 else 'end'}"


@tool("docx_create", "Create a new .docx from a list of blocks: [{type:'heading'|'paragraph'|'bullet', text, level}].",
      {"path": S, "blocks": {"type": "array", "items": {"type": "object"}}}, ["path", "blocks"])
def docx_create(cfg, path, blocks):
    p = resolve(cfg, path, write=True)
    d = _docx().Document()
    for b in blocks:
        t = b.get("type", "paragraph")
        if t == "heading":
            d.add_heading(b["text"], level=b.get("level", 1))
        elif t == "bullet":
            d.add_paragraph(b["text"], style="List Bullet")
        else:
            d.add_paragraph(b["text"])
    p.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(p))
    return f"created {p}"


def _xl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise ValueError("openpyxl missing: pip install openpyxl")


@tool("xlsx_read", "Read an .xlsx sheet (or range like 'A1:D20').", {"path": S, "sheet": S, "range": S}, ["path"])
def xlsx_read(cfg, path, sheet=None, range=None):
    wb = _xl().load_workbook(str(resolve(cfg, path)))
    if sheet:
        if sheet not in wb.sheetnames:
            raise ValueError(f"sheet '{sheet}' not found (available: {wb.sheetnames})")
        ws = wb[sheet]
    else:
        ws = wb.active
    rows = ws[range] if range else ws.iter_rows()
    out = [f"sheets: {wb.sheetnames}"]
    for r in rows:
        out.append(" | ".join("" if c.value is None else str(c.value) for c in r))
    return "\n".join(out[:300])


@tool("xlsx_set_cell", "Set a cell value or formula (e.g. '=SUM(A1:A9)').",
      {"path": S, "sheet": S, "cell": S, "value": S}, ["path", "cell", "value"])
def xlsx_set_cell(cfg, path, cell, value, sheet=None):
    p = resolve(cfg, path, write=True)
    wb = _xl().load_workbook(str(p))
    if sheet:
        if sheet not in wb.sheetnames:
            raise ValueError(f"sheet '{sheet}' not found (available: {wb.sheetnames})")
        ws = wb[sheet]
    else:
        ws = wb.active
    ws[cell] = value
    wb.save(str(p))
    return f"{cell} = {value}"


@tool("pptx_read", "Read all slide text from a .pptx.", {"path": S}, ["path"])
def pptx_read(cfg, path):
    from pptx import Presentation
    pr = Presentation(str(resolve(cfg, path)))
    out = []
    for i, s in enumerate(pr.slides, 1):
        out.append(f"--- slide {i} ---")
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
    return "\n".join(out)